from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_
from typing import Optional
import io, csv
from datetime import datetime

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.server import Server, CustomField

router = APIRouter()

HEADERS = [
    "SL","NAME","PRIVATE IP","BUSINESS APP","HOSTNAME",
    "OS","VERSION","TYPE","ENV","PLATFORM","OPS MANAGER",
    "AZURE UPDATE SCHEDULE","DEFENDER","PAM","IPA INTEGRATION",
    "MSB COMPLIANCE","SUBSCRIPTION","LOCATION","RESOURCE GROUP",
    "OWNER","OWNER TEAM","CRITICALITY","SUPPORT SUMA",
    "SUPPORT DCENTRAL","DCENTRAL","CROWDSTRIKE","STATUS","NOTES"
]

def apply_filters(q, search, os_name, os_version, status, environment, platform, location, owner_team, criticality, subscription, crowdstrike, onboarded_defender):
    if search:
        q = q.filter(or_(
            Server.hostname.ilike(f"%{search}%"), Server.name.ilike(f"%{search}%"),
            Server.private_ip.ilike(f"%{search}%"), Server.business_app.ilike(f"%{search}%"),
            Server.owner.ilike(f"%{search}%"), Server.owner_team.ilike(f"%{search}%"),
        ))
    if os_name:            q = q.filter(Server.os_name.ilike(f"%{os_name}%"))
    if os_version:         q = q.filter(Server.os_version.ilike(f"%{os_version}%"))
    if status:             q = q.filter(Server.status.ilike(f"%{status}%"))
    if environment:        q = q.filter(Server.environment.ilike(f"%{environment}%"))
    if platform:           q = q.filter(Server.platform.ilike(f"%{platform}%"))
    if location:           q = q.filter(Server.location.ilike(f"%{location}%"))
    if owner_team:         q = q.filter(Server.owner_team.ilike(f"%{owner_team}%"))
    if criticality:        q = q.filter(Server.criticality.ilike(f"%{criticality}%"))
    if subscription:       q = q.filter(Server.subscription.ilike(f"%{subscription}%"))
    if crowdstrike:        q = q.filter(Server.crowdstrike.ilike(f"%{crowdstrike}%"))
    if onboarded_defender: q = q.filter(Server.onboarded_defender.ilike(f"%{onboarded_defender}%"))
    return q

def get_servers(db, **kw):
    q = apply_filters(db.query(Server), **kw)
    return q.order_by(Server.sl, Server.id).all()

def row_vals(s, custom_fields):
    row = [
        s.sl, s.name, s.private_ip, s.business_app, s.hostname,
        s.os_name, s.os_version, s.server_type, s.environment, s.platform,
        s.ops_manager, s.azure_update_schedule, s.onboarded_defender, s.onboarded_pam,
        s.ipa_integration, s.msb_compliance, s.subscription, s.location, s.resource_group,
        s.owner, s.owner_team, s.criticality, s.support_suma, s.support_dcentral,
        s.dcentral, s.crowdstrike, s.status, s.notes or ""
    ]
    cd = s.custom_data or {}
    for cf in custom_fields:
        row.append(cd.get(cf.name, ""))
    return row

FILTER_PARAMS = ["search","os_name","os_version","status","environment","platform",
                 "location","owner_team","criticality","subscription","crowdstrike","onboarded_defender"]

def get_filter_deps():
    from fastapi import Query as Q
    return {k: None for k in FILTER_PARAMS}

@router.get("/csv")
def export_csv(
    search: Optional[str]=None, os_name: Optional[str]=None, os_version: Optional[str]=None,
    status: Optional[str]=None, environment: Optional[str]=None, platform: Optional[str]=None,
    location: Optional[str]=None, owner_team: Optional[str]=None, criticality: Optional[str]=None,
    subscription: Optional[str]=None, crowdstrike: Optional[str]=None, onboarded_defender: Optional[str]=None,
    db: Session=Depends(get_db), current_user=Depends(get_current_user)
):
    cf = db.query(CustomField).all()
    servers = get_servers(db, search=search, os_name=os_name, os_version=os_version, status=status,
        environment=environment, platform=platform, location=location, owner_team=owner_team,
        criticality=criticality, subscription=subscription, crowdstrike=crowdstrike, onboarded_defender=onboarded_defender)
    headers = HEADERS + [c.label for c in cf]
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(headers)
    for s in servers:
        w.writerow(row_vals(s, cf))
    out.seek(0)
    fname = f"server_inventory_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return StreamingResponse(io.BytesIO(out.getvalue().encode()), media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={fname}"})

@router.get("/excel")
def export_excel(
    search: Optional[str]=None, os_name: Optional[str]=None, os_version: Optional[str]=None,
    status: Optional[str]=None, environment: Optional[str]=None, platform: Optional[str]=None,
    location: Optional[str]=None, owner_team: Optional[str]=None, criticality: Optional[str]=None,
    subscription: Optional[str]=None, crowdstrike: Optional[str]=None, onboarded_defender: Optional[str]=None,
    db: Session=Depends(get_db), current_user=Depends(get_current_user)
):
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    cf = db.query(CustomField).all()
    servers = get_servers(db, search=search, os_name=os_name, os_version=os_version, status=status,
        environment=environment, platform=platform, location=location, owner_team=owner_team,
        criticality=criticality, subscription=subscription, crowdstrike=crowdstrike, onboarded_defender=onboarded_defender)
    headers = HEADERS + [c.label for c in cf]

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Server Inventory"
    hfill = PatternFill("solid", fgColor="1E3A5F")
    hfont = Font(bold=True, color="FFFFFF", size=10)
    status_colors = {"active":"E8F5E9","inactive":"FFF8E1","to be decommissioned":"FFEBEE"}

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = hfill; cell.font = hfont
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = max(len(h)+3, 12)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for ri, s in enumerate(servers, 2):
        sc = status_colors.get((s.status or "").lower(), "FFFFFF")
        sfill = PatternFill("solid", fgColor=sc)
        for ci, val in enumerate(row_vals(s, cf), 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.fill = sfill
            cell.font = Font(size=10)

    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}1"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"server_inventory_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"})

@router.get("/pdf")
def export_pdf(
    search: Optional[str]=None, os_name: Optional[str]=None, os_version: Optional[str]=None,
    status: Optional[str]=None, environment: Optional[str]=None, platform: Optional[str]=None,
    location: Optional[str]=None, owner_team: Optional[str]=None, criticality: Optional[str]=None,
    subscription: Optional[str]=None, crowdstrike: Optional[str]=None, onboarded_defender: Optional[str]=None,
    db: Session=Depends(get_db), current_user=Depends(get_current_user)
):
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    servers = get_servers(db, search=search, os_name=os_name, os_version=os_version, status=status,
        environment=environment, platform=platform, location=location, owner_team=owner_team,
        criticality=criticality, subscription=subscription, crowdstrike=crowdstrike, onboarded_defender=onboarded_defender)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A3), leftMargin=10*mm, rightMargin=10*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=14, textColor=colors.HexColor("#1E3A5F"))
    story.append(Paragraph("Server Inventory Report", title_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Total: {len(servers)} servers", styles["Normal"]))
    story.append(Spacer(1, 6*mm))

    col_headers = ["SL","Hostname","Private IP","OS","Ver","Type","Env","Platform","Location","Subscription","Owner","Team","Criticality","Defender","CrowdStrike","Status"]
    data = [col_headers]
    sc_map = {"active": colors.HexColor("#E8F5E9"), "inactive": colors.HexColor("#FFF8E1"), "to be decommissioned": colors.HexColor("#FFEBEE")}
    row_colors = []
    for i, s in enumerate(servers):
        data.append([s.sl or "", s.hostname, s.private_ip or "", s.os_name or "",
                     s.os_version or "", s.server_type or "", s.environment or "",
                     s.platform or "", s.location or "", s.subscription or "",
                     s.owner or "", s.owner_team or "", s.criticality or "",
                     s.onboarded_defender or "", s.crowdstrike or "", s.status or ""])
        row_colors.append((i+1, sc_map.get((s.status or "").lower(), colors.white)))

    col_w = [10,35,28,20,16,22,18,22,22,28,30,26,20,18,22,28]
    col_w = [w*mm for w in col_w]
    t = Table(data, colWidths=col_w, repeatRows=1)
    cmds = [
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),7),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F5F5F5")]),
        ("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#CCCCCC")),
        ("ALIGN",(0,0),(-1,-1),"LEFT"),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("PADDING",(0,0),(-1,-1),2),
    ]
    for ri, bg in row_colors:
        cmds.append(("BACKGROUND",(0,ri),(-1,ri),bg))
    t.setStyle(TableStyle(cmds))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    fname = f"server_inventory_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={fname}"})
